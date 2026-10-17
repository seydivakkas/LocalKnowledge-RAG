"""G03 full lifecycle synthetic fixtures; never include real customer documents."""
import io
import json
from pathlib import Path
import os
import sqlite3
import tempfile
import unittest
from contextlib import redirect_stdout, closing

from localknowledge.application import AppServices
from localknowledge.cli import main
from localknowledge.ingest.service import IntakeService
from localknowledge.ingest.parsers import IntakeError
from localknowledge.storage.db import open_db, DatabaseError


class G03IntegrationTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.base=Path(self.tmp.name)
        self.service=IntakeService(self.base/'data')
        self.source=self.base/'source'
        self.source.mkdir()

    def write(self,name,content):
        p=self.source/name
        p.write_bytes(content)
        return p

    def test_txt_import_counts_reopen_and_invariant(self):
        original='İzmir\nAnkara\nTürkçe kaynak\n'.encode()
        path=self.write('Belge.txt',original)
        r=self.service.import_document(path)
        self.assertEqual(path.read_bytes(),original)
        self.assertEqual(r.revision,1)
        self.assertEqual(r.operation,'CREATED')
        self.assertTrue(self.service.verify(r.document_id,1))
        after=IntakeService(self.base/'data')
        self.assertEqual(after.repository.stats().registered_documents,1)
        self.assertEqual(after.repository.stats().stored_versions,1)
        self.assertEqual(after.repository.stats().indexed_documents,0)
        self.assertEqual(after.repository.list_documents()[0].paragraph_count,3)
        self.assertEqual(after.repository.list_documents()[0].name,'Belge.txt')

    def test_duplicate_same_path_no_new_version(self):
        p=self.write('a.txt',b'abc\n')
        first=self.service.import_document(p)
        second=self.service.import_document(p)
        self.assertEqual(first.document_id,second.document_id)
        self.assertEqual(second.operation,'DUPLICATE')
        self.assertEqual(second.revision,1)
        self.assertEqual(self.service.repository.stats().stored_versions,1)

    def test_same_bytes_new_path_alias_no_new_document(self):
        a=self.write('a.txt',b'example')
        b=self.write('b.txt',b'example')
        x=self.service.import_document(a)
        y=self.service.import_document(b)
        self.assertEqual(y.operation,'ALIAS')
        self.assertEqual(x.document_id,y.document_id)
        self.assertEqual(self.service.repository.stats().registered_documents,1)

    def test_same_path_modified_creates_version(self):
        p=self.write('a.txt',b'v1')
        a=self.service.import_document(p)
        p.write_bytes(b'v2 contents')
        b=self.service.import_document(p)
        self.assertEqual(a.document_id,b.document_id)
        self.assertEqual(b.revision,2)
        self.assertTrue(self.service.verify(a.document_id,1))
        self.assertTrue(self.service.verify(a.document_id,2))
        self.assertEqual(len(self.service.repository.versions(a.document_id)),2)

    def test_archive_is_reversible_and_non_destructive(self):
        p=self.write('a.txt',b'abc')
        doc=self.service.import_document(p)
        self.service.repository.archive(doc.document_id)
        self.assertEqual(self.service.repository.stats().registered_documents,0)
        self.assertEqual(self.service.repository.list_documents(),())
        self.assertEqual(len(self.service.repository.list_documents(include_archived=True)),1)
        self.assertTrue(self.service.verify(doc.document_id,1))
        self.service.repository.archive(doc.document_id,False)
        self.assertEqual(self.service.repository.stats().registered_documents,1)

    def test_filter_and_safe_literal_like(self):
        self.service.import_document(self.write('100%.txt',b'p'))
        self.service.import_document(self.write('beta_.txt',b'b'))
        self.assertEqual([r.name for r in self.service.repository.list_documents('%')],['100%.txt'])
        self.assertEqual([r.name for r in self.service.repository.list_documents('_')],['beta_.txt'])
        self.assertEqual(self.service.repository.count_documents(kind='PDF'),0)

    def test_invalid_magic_does_not_create_doc(self):
        p=self.write('fake.pdf',b'not a PDF at all')
        with self.assertRaises(IntakeError) as error:self.service.import_document(p)
        self.assertEqual(error.exception.code,'WRONG_FORMAT')
        self.assertEqual(self.service.repository.stats().registered_documents,0)

    def test_binary_txt_rejected(self):
        with self.assertRaises(IntakeError) as error:self.service.import_document(self.write('x.txt',b'\x00\x01\x02'))
        self.assertEqual(error.exception.code,'BINARY_TEXT')

    def test_invalid_docx_rejected(self):
        with self.assertRaises(IntakeError) as error:self.service.import_document(self.write('x.docx',b'broken'))
        self.assertEqual(error.exception.code,'WRONG_FORMAT')

    def test_pdf_pages_have_real_numbers_not_search_chunks(self):
        from reportlab.pdfgen.canvas import Canvas
        p=self.source/'test.pdf'
        c=Canvas(str(p));c.drawString(90,700,'Test page 1');c.showPage();c.drawString(90,700,'Test page 2');c.save()
        r=self.service.import_document(p)
        self.assertEqual(self.service.repository.list_documents()[0].page_count,2)
        with closing(sqlite3.connect(self.service.db_path)) as conn:
            pages=conn.execute('SELECT page_number FROM page_metadata').fetchall()
            self.assertEqual(pages,[(1,),(2,)])
            self.assertEqual(conn.execute('PRAGMA user_version').fetchone()[0],3)
        self.assertTrue(self.service.verify(r.document_id,1))

    def test_docx_real_paragraph_anchors_without_page(self):
        from docx import Document
        p=self.source/'word.docx'
        d=Document();d.add_paragraph('İlk paragraf');d.add_paragraph('İkinci paragraf');d.save(p)
        r=self.service.import_document(p)
        self.assertEqual(self.service.repository.list_documents()[0].page_count,None)
        self.assertEqual(self.service.repository.list_documents()[0].paragraph_count,2)
        with closing(sqlite3.connect(self.service.db_path)) as conn:
            data=conn.execute('SELECT ordinal,char_start,char_end FROM paragraph_metadata ORDER BY ordinal').fetchall()
        self.assertEqual(len(data),2)
        self.assertTrue(all(e>s for _,s,e in data))

    def test_corrupted_original_object_fails_closed(self):
        r=self.service.import_document(self.write('a.txt',b'good'))
        rel,_,_=self.service.repository.managed_object(r.document_id,1)
        object_path=self.service.root/rel
        object_path.write_bytes(b'evil')
        self.assertFalse(self.service.verify(r.document_id,1))
        with self.assertRaises(IntakeError) as error:self.service.import_document(self.write('b.txt',b'good'))
        self.assertEqual(error.exception.code,'OBJECT_CORRUPT')

    def test_unknown_schema_fails_without_mutation(self):
        db=self.base/'unsupported.sqlite'
        with closing(sqlite3.connect(db)) as conn:
            conn.execute('PRAGMA user_version=99')
        before=db.read_bytes()
        with self.assertRaises(DatabaseError):open_db(db)
        self.assertEqual(before,db.read_bytes())

    def test_unknown_existing_unversioned_schema_rejected(self):
        db=self.base/'bad.sqlite'
        with closing(sqlite3.connect(db)) as conn:conn.execute('CREATE TABLE old_customer(id INTEGER)')
        with self.assertRaises(DatabaseError):open_db(db)
        with closing(sqlite3.connect(db)) as conn:self.assertEqual(conn.execute('SELECT name FROM sqlite_master WHERE name="old_customer"').fetchone()[0],'old_customer')

    def test_cli_import_list_stats(self):
        path=self.write('a.txt',b'hello offline')
        for argv in [['--import-file',str(path)],['--list-documents'],['--library-stats']]:
            output=io.StringIO()
            with redirect_stdout(output):self.assertEqual(main([*argv,'--data-dir',str(self.service.root)]),0)
            value=json.loads(output.getvalue())
        self.assertEqual(value['registered_documents'],1)
        self.assertEqual(value['indexed_documents'],0)

    def test_sqlite_integrity_and_import_events(self):
        p=self.write('a.txt',b'example')
        self.service.import_document(p)
        self.service.import_document(p)
        with closing(sqlite3.connect(self.service.db_path)) as conn:
            self.assertEqual(conn.execute('PRAGMA integrity_check').fetchone()[0],'ok')
            self.assertEqual(conn.execute('PRAGMA foreign_key_check').fetchall(),[])
            self.assertEqual([x[0] for x in conn.execute('SELECT operation FROM import_events ORDER BY id')],["CREATED","DUPLICATE"])


if __name__=='__main__': unittest.main()
