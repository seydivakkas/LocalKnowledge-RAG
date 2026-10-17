"""G05: FTS5/BM25, current version, source SHA and absolute excerpt offsets."""
import hashlib
import sqlite3
import tempfile
from contextlib import closing
import unittest
from pathlib import Path
from localknowledge.ingest.service import IntakeService
from localknowledge.jobs.runner import Runner
from localknowledge.search.service import SearchService, safe_query, SearchError
from localknowledge.search.index import split_unit, rebuild_completed
from localknowledge.storage.db import connect, transaction, V2, SCHEMA, open_db

class G05SearchTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.base=Path(self.temp.name)
        self.root=self.base/'data';self.ingest=IntakeService(self.root)
        self.searcher=SearchService(self.root)
    def tearDown(self):self.temp.cleanup()
    def write(self,name,body):
        path=self.base/name;path.write_text(body,encoding='utf-8');return path
    def process(self):
        runner=Runner(self.root,poll_interval=.005)
        self.assertTrue(runner.run_one())
    def test_fts_keyword_bm25_and_verified_excerpt(self):
        res=self.ingest.import_document(self.write('evrak.txt','Koridor genişliği yüz santim olmalıdır.\nDüzlem çizimi ve cephe analizi.\n'))
        self.assertEqual(self.searcher.stats(),{'indexed_documents':0,'indexed_chunks':0})
        self.process()
        self.assertEqual(self.searcher.stats(),{'indexed_documents':1,'indexed_chunks':2})
        hits=self.searcher.search('koridor genişliği')
        self.assertEqual(len(hits),1)
        hit=hits[0]
        self.assertEqual((hit.document_id,hit.revision,hit.unit_number,hit.char_start),(res.document_id,1,0,0))
        self.assertTrue(hit.verified)
        self.assertLess(hit.bm25_score,0)
        self.assertEqual(hit.source_label.split('paragraf')[1].split(',')[0].strip(),'1')
        detail=self.searcher.source_passage(hit.chunk_id)
        self.assertIsNotNone(detail)
        self.assertEqual(detail[1][detail[2]:detail[3]],hit.exact_quote)
        self.assertEqual(self.searcher.search('yokkelime'),())

    def test_revision_update_archive_drops_index_and_restore_must_reindex(self):
        path=self.write('rev.txt','Eski cephe planı.\n')
        first=self.ingest.import_document(path);self.process()
        old=self.searcher.search('eski')[0]
        path.write_text('Yeni yangın merdiveni raporu.\n',encoding='utf-8')
        new=self.ingest.import_document(path)
        self.assertEqual(new.revision,2)
        self.assertEqual(self.searcher.search('eski'),())
        self.assertIsNone(self.searcher.resolve(old.chunk_id))
        self.process()
        self.assertEqual(len(self.searcher.search('yangın')),1)
        self.ingest.repository.archive(new.document_id,True)
        self.assertEqual(self.searcher.search('yangın'),())
        self.assertEqual(self.searcher.stats()['indexed_documents'],0)
        self.ingest.repository.archive(new.document_id,False)
        # Restore must re-enable an already completed G04 extraction.
        self.assertEqual(len(self.searcher.search('yangın')),1)

    def test_original_corruption_hides_hits_and_sources(self):
        res=self.ingest.import_document(self.write('sakla.txt','Korunan metin Türkçe.\n'))
        self.process();hit=self.searcher.search('korunan')[0]
        rel,_,_=self.ingest.repository.managed_object(res.document_id,1)
        (self.root/rel).write_bytes(b'X'*(self.root/rel).stat().st_size)
        self.assertEqual(self.searcher.search('korunan'),())
        self.assertIsNone(self.searcher.resolve(hit.chunk_id))
        self.assertIsNone(self.searcher.source_passage(hit.chunk_id))

    def test_fts5_operator_injection_cannot_match_unrelated_terms(self):
        self.ingest.import_document(self.write('a.txt','Mimari projede koridor alanı.\n'))
        self.process()
        for q in ['" OR gizli','koridor* NEAR/3 test', 'koridor OR alınmaz', "koridor' OR 1=1 --"]:
            self.assertIsInstance(self.searcher.search(q),tuple)
        for q in ['', ' * ', 'X'*250, 'çok '*17]:
            with self.assertRaises(SearchError):self.searcher.search(q)
        self.assertEqual(self.searcher.search('koridor',kind='PDF'),())
        self.assertEqual(len(self.searcher.search('koridor',kind='TXT')),1)
        self.assertEqual(self.searcher.search('koridor',created_from='2099-01-01'),())
        with self.assertRaises(SearchError):self.searcher.search('koridor',created_from='2026; DROP TABLE')

    def test_exact_chunk_boundaries_and_unicode_source_preserved(self):
        original=('İğne, ışık, Çizim — "+ özel" açıklamaları. '*30).strip()
        chunks=list(split_unit(original,limit=80,overlap=12))
        self.assertGreater(len(chunks),10)
        self.assertTrue(all(original[a:b]==chunk for a,b,chunk in chunks))
        self.assertEqual(chunks[0][0],0)
        self.assertEqual(chunks[-1][1],len(original))

    def test_pdf_physical_page_and_docx_paragraph_source(self):
        from reportlab.pdfgen.canvas import Canvas
        from docx import Document
        pdf=self.base/'sayfa.pdf';cv=Canvas(str(pdf));cv.drawString(50,700,'Cephe tasarimi birinci sayfa')
        cv.showPage();cv.drawString(50,700,'Koridor genişligi ikinci sayfa');cv.save()
        rpdf=self.ingest.import_document(pdf)
        word=self.base/'kriter.docx';d=Document();d.add_paragraph('Teknik uygulama');d.add_paragraph('Sürdürülebilir yapılar tasarım yaklaşımı');d.save(word)
        rdocx=self.ingest.import_document(word)
        self.process();self.process()
        pdf_hit=self.searcher.search('koridor',kind='PDF')[0]
        self.assertEqual((pdf_hit.unit_kind,pdf_hit.unit_number),( 'PAGE',1))
        self.assertIn('sayfa 2',pdf_hit.source_label)
        docx_hit=self.searcher.search('Sürdürülebilir',kind='DOCX')[0]
        self.assertEqual(docx_hit.unit_kind,'PARAGRAPH')
        self.assertEqual(docx_hit.unit_number,1)
        self.assertEqual(docx_hit.document_id,rdocx.document_id)
        self.assertEqual(len(self.searcher.search('koridor',kind='DOCX')),0)

    def test_search_index_rebuild_is_transactional_and_idempotent(self):
        self.ingest.import_document(self.write('yeniden.txt','Sistem arama kaynakları.\n'));self.process()
        with connect(self.ingest.db_path) as conn,transaction(conn):
            report=rebuild_completed(conn)
        self.assertEqual(report,{'documents':1,'chunks':1})
        self.assertEqual(len(self.searcher.search('arama')),1)
        with connect(self.ingest.db_path) as conn,transaction(conn):
            report2=rebuild_completed(conn)
        self.assertEqual(report2,report)
        self.assertEqual(len(self.searcher.search('arama')),1)

    def test_fts_delete_trigger_prevents_stale_ghost_search(self):
        r=self.ingest.import_document(self.write('silin.txt','Güncel şema kontrolü.\n'));self.process()
        with connect(self.ingest.db_path) as conn,transaction(conn):
            conn.execute('DELETE FROM search_chunks WHERE document_id=?',(r.document_id,))
            conn.execute('DELETE FROM indexed_versions WHERE document_id=?',(r.document_id,))
        self.assertEqual(self.searcher.search('şema'),())

    def test_v2_completed_extraction_backfill_into_v3(self):
        db=self.base/'legacy_v2.sqlite'
        with closing(sqlite3.connect(db)) as raw:
            raw.executescript(SCHEMA+V2)
            raw.execute('PRAGMA user_version=2')
            raw.execute("INSERT INTO documents(id,name,kind) VALUES('a','legacy.txt','TXT')")
            raw.execute("INSERT INTO versions(document_id,revision,sha256,byte_size,object_relpath,paragraph_count,original_name) VALUES('a',1,?,6,?,1,'legacy.txt')",('b'*64,'objects/bb/'+('b'*64)+'.bin'))
            raw.execute("INSERT INTO index_jobs(document_id,revision,sha256,status,total_units,processed_units) VALUES('a',1,?,'COMPLETED',1,1)",('b'*64,))
            job_id=raw.execute('SELECT id FROM index_jobs').fetchone()[0]
            raw.execute("INSERT INTO extracted_spans(job_id,document_id,revision,unit_kind,unit_number,char_start,char_end,original_text) VALUES(?, 'a',1,'PARAGRAPH',0,0,6,'Legacy')",(job_id,))
            raw.commit()
        with connect(db) as conn:
            self.assertEqual(conn.execute('PRAGMA user_version').fetchone()[0],3)
            self.assertEqual(conn.execute('SELECT count(*) FROM search_chunks').fetchone()[0],1)
            self.assertEqual(conn.execute('SELECT count(*) FROM search_fts WHERE search_fts MATCH ?',('"Legacy"',)).fetchone()[0],1)
        with connect(db) as conn:self.assertEqual(conn.execute('SELECT count(*) FROM search_chunks').fetchone()[0],1)
